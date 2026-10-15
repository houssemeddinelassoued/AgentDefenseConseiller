"""Extraction du texte d'un document téléversé (PDF, Word, texte)."""

from __future__ import annotations

import io
import logging

logging.getLogger("pypdf").setLevel(logging.ERROR)  # avertissements de structure PDF non bloquants

FORMATS = ["pdf", "docx", "txt", "md"]
LIMITE_CARACTERES = 80_000  # ≈ 25 pages ; une demande type compte 4 à 5 pages


class DocumentIllisible(ValueError):
    """Le document ne peut pas être converti en texte exploitable."""


def extraire_texte(nom: str, donnees: bytes) -> str:
    extension = nom.rsplit(".", 1)[-1].lower() if "." in nom else ""
    if extension == "pdf":
        texte = _pdf(donnees)
    elif extension == "docx":
        texte = _docx(donnees)
    elif extension in ("txt", "md"):
        texte = _texte(donnees)
    else:
        raise DocumentIllisible(f"Format « .{extension} » non pris en charge (formats acceptés : {', '.join(FORMATS)}).")

    texte = "\n".join(ligne.rstrip() for ligne in texte.splitlines()).strip()
    if len(texte) < 200:
        raise DocumentIllisible(
            "Le document ne contient presque pas de texte extractible. S'il s'agit d'un PDF numérisé "
            "(image), copiez son contenu dans « Rédiger le besoin »."
        )
    if len(texte) > LIMITE_CARACTERES:
        # On refuse plutôt que de tronquer : une analyse sur un texte coupé serait trompeuse.
        raise DocumentIllisible(
            f"Le document est trop long ({len(texte):,} caractères, maximum {LIMITE_CARACTERES:,}). "
            "Soumettez la partie qui décrit le besoin (contexte, problème, résultats attendus, contraintes)."
        )
    return texte


def _pdf(donnees: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        lecteur = PdfReader(io.BytesIO(donnees))
        if lecteur.is_encrypted:
            raise DocumentIllisible("Le PDF est protégé par mot de passe.")
        return "\n\n".join(page.extract_text() or "" for page in lecteur.pages)
    except PdfReadError as exc:
        raise DocumentIllisible(f"PDF illisible : {exc}") from exc


def _docx(donnees: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(donnees))
    except Exception as exc:  # fichier corrompu ou pas un vrai .docx
        raise DocumentIllisible(f"Document Word illisible : {exc}") from exc
    blocs = [p.text for p in document.paragraphs]
    for tableau in document.tables:
        for ligne in tableau.rows:
            blocs.append(" | ".join(cellule.text.strip() for cellule in ligne.cells))
    return "\n".join(blocs)


def _texte(donnees: bytes) -> str:
    for encodage in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return donnees.decode(encodage)
        except UnicodeDecodeError:
            continue
    raise DocumentIllisible("Encodage du fichier texte non reconnu.")
